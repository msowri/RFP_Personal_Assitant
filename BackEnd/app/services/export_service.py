import logging
from typing import Dict, List
from docx import Document
from docx.shared import Inches
from io import BytesIO
import os

logger = logging.getLogger(__name__)

class ExportService:
    def __init__(self):
        pass

    def export_to_docx(self, content: Dict[str, any], filename: str = None) -> BytesIO:
        """Export content to DOCX format"""
        try:
            doc = Document()
            
            # Add title
            title = content.get('title', 'RFP Response Document')
            doc.add_heading(title, 0)
            
            # Add metadata
            if 'metadata' in content:
                doc.add_heading('Document Information', level=1)
                for key, value in content['metadata'].items():
                    doc.add_paragraph(f"{key}: {value}")
                doc.add_paragraph()  # Add space
            
            # Add Q&A pairs
            if 'qa_pairs' in content:
                doc.add_heading('Questions and Answers', level=1)
                for i, qa in enumerate(content['qa_pairs'], 1):
                    doc.add_heading(f"Q{i}: {qa.get('question', '')}", level=2)
                    doc.add_paragraph(qa.get('answer', ''))
                    doc.add_paragraph()  # Add space between Q&A
            
            # Add additional content
            if 'additional_content' in content:
                doc.add_heading('Additional Information', level=1)
                for section in content['additional_content']:
                    if 'title' in section:
                        doc.add_heading(section['title'], level=2)
                    if 'content' in section:
                        doc.add_paragraph(section['content'])
            
            # Save to BytesIO
            docx_buffer = BytesIO()
            doc.save(docx_buffer)
            docx_buffer.seek(0)
            
            logger.info(f"Successfully exported to DOCX format")
            return docx_buffer
            
        except Exception as e:
            logger.error(f"Error exporting to DOCX: {str(e)}")
            raise

    def export_drafts_to_docx(self, drafts: List[Dict], document_info: Dict = None) -> BytesIO:
        """Export multiple drafts to DOCX format"""
        try:
            doc = Document()
            
            # Add title
            doc.add_heading('RFP Response Drafts', 0)
            
            # Add document information if provided
            if document_info:
                doc.add_heading('Source Document', level=1)
                doc.add_paragraph(f"Document: {document_info.get('file_name', 'Unknown')}")
                doc.add_paragraph(f"Type: {document_info.get('file_type', 'Unknown')}")
                doc.add_paragraph(f"Uploaded: {document_info.get('created_on', 'Unknown')}")
                doc.add_paragraph()
            
            # Add drafts
            for i, draft in enumerate(drafts, 1):
                doc.add_heading(f'Draft {i}', level=1)
                
                if draft.get('question'):
                    doc.add_heading('Question:', level=2)
                    doc.add_paragraph(draft['question'])
                
                if draft.get('answer'):
                    doc.add_heading('Answer:', level=2)
                    doc.add_paragraph(draft['answer'])
                
                # Add metadata
                doc.add_paragraph(f"Status: {'Edited' if draft.get('is_edited', False) else 'Generated'}")
                doc.add_paragraph(f"Created: {draft.get('created_on', 'Unknown')}")
                doc.add_paragraph()  # Add space between drafts
            
            # Save to BytesIO
            docx_buffer = BytesIO()
            doc.save(docx_buffer)
            docx_buffer.seek(0)
            
            logger.info(f"Successfully exported {len(drafts)} drafts to DOCX format")
            return docx_buffer
            
        except Exception as e:
            logger.error(f"Error exporting drafts to DOCX: {str(e)}")
            raise

    def export_to_text(self, content: Dict[str, any]) -> str:
        """Export content to plain text format"""
        try:
            lines = []
            
            # Add title
            title = content.get('title', 'RFP Response Document')
            lines.append("=" * len(title))
            lines.append(title)
            lines.append("=" * len(title))
            lines.append("")
            
            # Add metadata
            if 'metadata' in content:
                lines.append("DOCUMENT INFORMATION")
                lines.append("-" * 20)
                for key, value in content['metadata'].items():
                    lines.append(f"{key}: {value}")
                lines.append("")
            
            # Add Q&A pairs
            if 'qa_pairs' in content:
                lines.append("QUESTIONS AND ANSWERS")
                lines.append("-" * 25)
                for i, qa in enumerate(content['qa_pairs'], 1):
                    lines.append(f"Q{i}: {qa.get('question', '')}")
                    lines.append(f"A{i}: {qa.get('answer', '')}")
                    lines.append("")
            
            # Add additional content
            if 'additional_content' in content:
                lines.append("ADDITIONAL INFORMATION")
                lines.append("-" * 22)
                for section in content['additional_content']:
                    if 'title' in section:
                        lines.append(section['title'])
                        lines.append("-" * len(section['title']))
                    if 'content' in section:
                        lines.append(section['content'])
                    lines.append("")
            
            return "\n".join(lines)
            
        except Exception as e:
            logger.error(f"Error exporting to text: {str(e)}")
            raise

    def get_supported_formats(self) -> Dict[str, List[str]]:
        """Get supported export formats"""
        return {
            "document_formats": ["docx", "txt"],
            "content_types": ["qa_pairs", "drafts", "full_document"],
            "features": [
                "Metadata preservation",
                "Q&A formatting", 
                "Multiple draft export",
                "Plain text fallback"
            ]
        }
